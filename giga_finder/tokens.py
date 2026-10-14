import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from giga_finder.argocd import open_pod_exec
from giga_finder.constants import *
from giga_finder.runtime import open_url
from giga_finder.search import iter_stand_configs
from giga_finder.ui import (
    action_label, colorize, input_with_timeout, is_help_request, print_action, print_divider,
    print_error, print_hint, print_info, print_warn,
)

def _import_openpyxl():
    """openpyxl для чтения/записи xlsx; если его нет в Python — ставим на лету (pip)."""
    try:
        import openpyxl
        return openpyxl
    except ImportError as error:
        # Диагностика: показываем, КАКИМ python запущен скрипт и что именно не нашлось.
        print_warn(f"openpyxl не найден ({error}).")
        print_info(f"Интерпретатор: {sys.executable}")
        print_info(f"Версия: {sys.version.split()[0]}")
        print_info(f"Пути поиска: {sys.path}")
    print_warn("Пробую установить через pip в этот же интерпретатор...")
    try:
        subprocess.run([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl"], check=True)
        import openpyxl
        print_info(f"openpyxl установлен: {openpyxl.__version__}")
        return openpyxl
    except Exception as error:
        print_error(f"Не удалось установить openpyxl автоматически: {error}")
        print_info("Установите вручную: python -m pip install openpyxl")
        return None
TOKENS_DIR_NAME = "tokens"
TOKENS_SOURCE_FILE = "tokens.xlsx"
def _create_blank_tokens_xlsx(path):
    """Создаёт пустой tokens.xlsx с заголовками: subdomain, ИНН."""
    openpyxl = _import_openpyxl()
    if openpyxl is None:
        raise RuntimeError("openpyxl недоступен")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "subdomain"
    ws["B1"] = "ИНН"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    wb.close()


def _maps_from_rows(rows):
    """subdomain → stand, число конфликтов, ИНН → stand, stand → зона ArgoCD.

    rows — один проход iter_stand_configs, чтобы не читать конфиги трижды.
    """
    subdomain = {}
    conflicts = 0
    inns = {}
    by_stand = {}
    code_by_label = {z["cluster"]: z["code"] for z in WORKSPACE_ZONES}
    for zone, _config_path, data in rows:
        stand = data.get("stand_name")
        sub = data.get("cowork_subdomain")
        if sub and stand:
            key = sub.strip().lower()
            if key in subdomain and subdomain[key] != stand:
                conflicts += 1
            else:
                subdomain[key] = stand
        inn = data.get("inn")
        if inn and stand:
            inn_clean = inn.strip().replace(chr(39), "").replace(chr(34), "")
            if inn_clean and inn_clean not in inns:
                inns[inn_clean] = stand
        if stand and stand not in by_stand:
            by_stand[stand] = {
                "code": code_by_label.get(zone.get("label", ""), ""),
                "app_template": zone.get("app_template"),
                "argocd_base": zone.get("argocd_base"),
                "admin_email": data.get("admin_email"),
            }
    return subdomain, conflicts, inns, by_stand


def _ch_num(value):
    value = value.strip()
    if value in ("", r"\N", "NULL"):
        return 0
    return int(float(value))


def _ch_split(line):
    """(kind, head, nums). kind — tab или space. nums — четыре хвостовых поля."""
    parts = line.split("\t")
    if len(parts) >= 5:
        return "tab", parts[0], parts[1:5]
    toks = line.split()
    if len(toks) < 5:
        return None
    return "space", toks, toks[-4:]


def _ch_metrics(nums):
    try:
        calls, inp, outp, tot = (_ch_num(n) for n in nums)
    except ValueError:
        return None
    if tot == 0:
        tot = inp + outp
    return calls, inp, outp, tot


def _token_utilization_sql(stands, days):
    """Один запрос на все стенды: GROUP BY environment, чтобы результат раскидать по строкам."""
    quoted = ", ".join(f"'{s}'" for s in stands)
    return (
        "SELECT environment, count() AS calls, "
        "sum(usage_details['input']) AS input_tokens, sum(usage_details['output']) AS output_tokens, "
        "sum(usage_details['total']) AS total_tokens "
        "FROM default.observations "
        f"WHERE environment IN ({quoted}) AND type = 'GENERATION' "
        f"AND start_time >= now() - INTERVAL {days} DAY "
        "GROUP BY environment FORMAT TabSeparated"
    )
def _parse_ch_env_output(raw_lines):
    """Разбор вывода clickhouse-client с колонкой environment.
    Возвращает dict: environment → (calls, inp, outp, tot)."""
    result = {}
    for line in raw_lines:
        split = _ch_split(line)
        if split is None:
            continue
        kind, head, nums = split
        metrics = _ch_metrics(nums)
        if metrics is None:
            continue
        env_name = head if kind == "tab" else head[0]
        result[env_name.strip()] = metrics
    return result
TOKENS_RESULT_FILE = "результат clickhouse.txt"
TOKENS_USERS_RESULT_FILE = "результат пользователи.txt"

_LIST_USERS_HEREDOC = """/app/services/auth/.venv/bin/python <<'PY'
import asyncio, sys, httpx
from datetime import UTC, datetime
try:
    from auth_service.bootstrap.config import get_settings
except ModuleNotFoundError:
    from auth_service.config import get_settings
try:
    from sqlalchemy import select
    from auth_sdk import AuthClient, KeycloakProviderConfig
except ModuleNotFoundError:
    from heimdall_auth import AuthClient, KeycloakProviderConfig
    from sqlalchemy import select
try:
    from auth_service.infrastructure.db.client import Database
except ModuleNotFoundError:
    from auth_service.db.client import Database
try:
    from auth_service.infrastructure.db.tables import InviteTokenModel
except ModuleNotFoundError:
    try:
        from auth_service.db.tables import InviteTokenModel
    except ModuleNotFoundError:
        from auth_service.db.migration_models import InviteTokenModel

ADMIN_EMAIL = __ADMIN_EMAIL__

async def main():
    settings = get_settings()
    base_url = settings.provider_issuer_url.split("/realms/")[0]
    realm = settings.provider_realm

    async with httpx.AsyncClient() as http:
        token_resp = await http.post(
            f"{base_url}/realms/{realm}/protocol/openid-connect/token",
            data={
                "grant_type": "client_credentials",
                "client_id": settings.provider_client_id,
                "client_secret": settings.provider_client_secret,
            },
        )
        token_resp.raise_for_status()
        access_token = token_resp.json()["access_token"]

        all_users = []
        first = 0
        page_size = 100
        while True:
            resp = await http.get(
                f"{base_url}/admin/realms/{realm}/users",
                params={"first": first, "max": page_size},
                headers={"Authorization": f"Bearer {access_token}"},
            )
            resp.raise_for_status()
            users = resp.json()
            if not users:
                break
            all_users.extend(users)
            if len(users) < page_size:
                break
            first += page_size

        print(f"Total: {len(all_users)}")
        for u in sorted(all_users, key=lambda x: x.get("email", "").lower()):
            email = u.get("email", "\u2014")
            name = f"{u.get('firstName', '')} {u.get('lastName', '')}".strip() or "\u2014"
            uid = u.get("id", "\u2014")
            status = "active" if u.get("enabled") else "disabled"
            ts = u.get("createdTimestamp")
            created = datetime.fromtimestamp(ts / 1000, tz=UTC).strftime("%Y-%m-%d") if ts else "\u2014"
            print(f"{email}  |  {name}  |  {status}  |  created={created}  |  id={uid}")

        # Проверка приглашения админа (http ещё жив внутри async with)
        if ADMIN_EMAIL:
            db = Database(settings)
            try:
                auth_client = AuthClient(
                    config=KeycloakProviderConfig(
                        issuer_url=settings.provider_issuer_url,
                        realm=settings.provider_realm,
                        client_id=settings.provider_client_id,
                        client_secret=settings.provider_client_secret,
                    ),
                    http_client=http,
                )
                user = await auth_client.find_user_by_email(email=ADMIN_EMAIL)
                if user is not None:
                    print(f"INVITE_STATUS: ACCEPTED {ADMIN_EMAIL}")
                else:
                    async with db.session() as session:
                        result = await session.execute(
                            select(InviteTokenModel)
                            .where(InviteTokenModel.email_normalized == ADMIN_EMAIL)
                            .order_by(InviteTokenModel.updated_at.desc())
                        )
                        tok = result.scalars().first()
                    if tok is None:
                        print(f"INVITE_STATUS: UNINVITED {ADMIN_EMAIL}")
                    else:
                        exp = tok.expires_at
                        if exp.tzinfo is None:
                            exp = exp.replace(tzinfo=UTC)
                        if tok.status == "pending" and exp > datetime.now(UTC):
                            print(f"INVITE_STATUS: WAITING {ADMIN_EMAIL}")
                        elif exp <= datetime.now(UTC):
                            print(f"INVITE_STATUS: EXPIRED {ADMIN_EMAIL}")
                        else:
                            print(f"INVITE_STATUS: OTHER {ADMIN_EMAIL}")
            finally:
                await db.close()

asyncio.run(main())
PY"""


def generate_check_invites_script(admin_email):
    """Скрипт: список пользователей + статус приглашения админа."""
    return _LIST_USERS_HEREDOC.replace("__ADMIN_EMAIL__", repr(admin_email))


def _parse_users_output(raw_text):
    """Из вывода скрипта извлекает число пользователей и статус приглашения.
    Возвращает (count: int|None, invite_status: str|None)."""
    clean = re.sub(r"\x1b\[[0-9;]*m", "", raw_text)
    m = re.search(r"Total:\s*(\d+)", clean)
    count = int(m.group(1)) if m else None
    im = re.search(r"INVITE_STATUS:\s*\w+", clean)
    invite = im.group(0).split(":")[1].strip() if im else None
    return count, invite


def _count_users_for_stands(stands, stand_to_zone, tokens_dir):
    """По очереди открывает heimdall-под для каждого стенда, принимает результат.
    Один скрипт на стенд: список пользователей + статус приглашения админа.
    Возвращает dict: stand_name → (int, str|None)."""
    results = {}
    result_path = tokens_dir / TOKENS_USERS_RESULT_FILE
    EXEC_DIR.mkdir(parents=True, exist_ok=True)
    for idx, stand in enumerate(stands, 1):
        zone = stand_to_zone.get(stand)
        if zone is None:
            print_warn(f"Зона не найдена для стенда {stand} — пропускаю.")
            continue
        admin_email = (zone.get("admin_email") or "").strip()
        print_divider()
        print_info(f"Стенд {idx}/{len(stands)}: {stand}")
        open_pod_exec(zone["code"], stand, "heimdall", zone.get("app_template"), zone.get("argocd_base"))
        # Генерируем скрипт с конкретным admin_email для этого стенда.
        script_content = generate_check_invites_script(admin_email)
        script_path = EXEC_DIR / "list_users.sh"
        script_path.write_text(script_content, encoding="utf-8")
        print_info(f"Скопируйте содержимое {script_path.name} целиком в веб-терминал пода и запустите.")
        result_path.write_text("", encoding="utf-8")
        print_info(f"Скопируйте вывод команды в файл: {result_path}")
        answer = input(f"{action_label()} Заполнили файл? (yes/no/skip, Enter — yes): ").strip().lower()
        if answer == "skip":
            print_warn(f"Пропущен стенд {stand}.")
            continue
        if answer in ("n", "no"):
            print_warn("Операция отменена.")
            break
        try:
            raw = result_path.read_text(encoding="utf-8")
        except OSError as error:
            print_error(f"Не удалось прочитать файл: {error}")
            continue
        count, invite = _parse_users_output(raw)
        if count is None:
            print_warn(f"Не удалось разобрать результат для {stand}: строка «Total: N» не найдена.")
            answer2 = input_with_timeout(f"{action_label()} Повторить? (yes/no/skip): ", INPUT_TIMEOUT)
            if answer2 and answer2.strip().lower() in ("y", "yes"):
                continue
            print_warn(f"Пропущен стенд {stand}.")
            continue
        results[stand] = (count, invite)
        invite_label = f", приглашение: {invite}" if invite else ""
        print_info(f"Пользователей на {stand}: {_fmt_int(count)}{invite_label}")
    return results

def _read_result_file(result_path):
    """Строки результата из промежуточного файла; пропуская служебные строки
    веб-терминала (root@..., #, $, эхо команды) и пустые."""
    try:
        raw = result_path.read_text(encoding="utf-8")
    except OSError as error:
        print_error(f"Не удалось прочитать файл: {error}")
        return []
    lines = []
    for line in raw.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith(("root@", "#", "$")) or "clickhouse-client" in stripped:
            continue
        lines.append(line.rstrip())
    return lines
def _write_token_report_xlsx(src_path, out_path, row_stands, totals_by_env, days, users_by_env=None, invite_by_env=None):
    """Копирует исходную таблицу и заполняет столбец D аналитикой по токенам.
    row_stands — список stand_name (или None), выровненный по видимым непустым строкам.
    Обрабатывает ТОЛЬКО видимые (развернутые) строки; свернутые (hidden) — не трогает."""
    openpyxl = _import_openpyxl()
    if openpyxl is None:
        raise RuntimeError("openpyxl недоступен")
    wb = openpyxl.load_workbook(str(src_path))
    ws = wb[wb.sheetnames[0]]
    filled = no_data = missing = skipped_hidden = 0
    idx = 0
    for row in range(2, ws.max_row + 1):
        # Свернутая строка — пропускаем целиком, ячейку D не трогаем.
        if ws.row_dimensions[row].hidden:
            skipped_hidden += 1
            continue
        cell_a = ws.cell(row=row, column=1).value
        cell_b = ws.cell(row=row, column=2).value
        a_empty = cell_a is None or str(cell_a).strip() == ""
        b_empty = cell_b is None or str(cell_b).strip() == ""
        if a_empty and b_empty:
            continue
        stand = row_stands[idx] if idx < len(row_stands) else None
        idx += 1
        if stand is None:
            text = "Не найден в конфигурациях репозитория"
            missing += 1
        elif stand not in totals_by_env:
            text = f"Данных нет: за последние {days} дн. вызовов LLM не найдено"
            no_data += 1
        else:
            calls, inp, outp, tot = totals_by_env[stand]
            parts = [f"Вызовов LLM: {_fmt_int(calls)}"]
            if users_by_env and stand in users_by_env:
                parts.append(f"Пользователей: {_fmt_int(users_by_env[stand])}")
            if invite_by_env and stand in invite_by_env:
                parts.append(f"Приглашение админа: {invite_by_env[stand]}")
            parts.append(f"Потрачено за период: {_fmt_int(tot)} токенов")
            parts.append(f"  запрос к модели (input):   {_fmt_int(inp)}")
            parts.append(f"  ответ модели (output):     {_fmt_int(outp)}")
            text = "\n".join(parts)
            filled += 1
        cell = ws.cell(row=row, column=4, value=text)
        cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    wb.save(str(out_path))
    return filled, no_data, missing, skipped_hidden
def token_utilization_report(repo_root, script_dir, zones):
    """Подсчёт токенов: tokens.xlsx → один запрос ClickHouse → столбец D нового xlsx."""
    print_divider()
    print_action("Подсчёт утилизации токенов")
    tokens_dir = script_dir / TOKENS_DIR_NAME
    if not tokens_dir.exists():
        tokens_dir.mkdir(parents=True)
        print_info(f"Создана папка: {tokens_dir}")
    else:
        print_info(f"Папка готова: {tokens_dir}")
    # Исходная таблица: ищем в папке подсчёта, затем рядом со скриптом.
    src = tokens_dir / TOKENS_SOURCE_FILE
    if not src.exists():
        alt = script_dir / TOKENS_SOURCE_FILE
        if alt.exists():
            src = alt
        else:
            _create_blank_tokens_xlsx(src)
            print_info(f"Создан шаблон: {src} (столбец A — subdomain, столбец B — ИНН)")
    openpyxl = _import_openpyxl()
    if openpyxl is None:
        return
    try:
        # Без read_only: нужен доступ к row_dimensions (hidden) — берём ТОЛЬКО
        # развернутые (видимые) строки, свернутые в запрос не включаем.
        wb = openpyxl.load_workbook(str(src))
        ws = wb[wb.sheetnames[0]]
        subs = []
        inns = []
        for row in range(2, ws.max_row + 1):
            if ws.row_dimensions[row].hidden:
                continue
            v_a = ws.cell(row=row, column=1).value
            v_b = ws.cell(row=row, column=2).value
            a = str(v_a).strip() if v_a is not None else ""
            b = str(v_b).strip() if v_b is not None else ""
            # Строка участвует, если заполнен хотя бы один критерий: subdomain (A) или ИНН (B).
            if not a and not b:
                continue
            subs.append(a)
            inns.append(b)
        wb.close()
    except Exception as error:
        print_error(f"Не удалось прочитать файл: {error}")
        return
    if not subs:
        print_error("В таблице нет ни одного subdomain (столбец A) или ИНН (столбец B).")
        return
    print_info(f"Видимых (развернутых) строк в таблице: {len(subs)}, уникальных subdomain: {len(set(s.lower() for s in subs))}.")
    days = prompt_period("Таймаут ожидания. Операция отменена.")
    if days is None:
        return
    print_info("Собираю маппинг из конфигураций репозитория...")
    mapping, conflicts, inn_mapping, stand_zones = _maps_from_rows(list(iter_stand_configs(repo_root, zones)))
    if conflicts:
        print_warn(f"Обратите внимание: {conflicts} конфликт(а) — один subdomain встречается в разных стендах; взят первый.")
    # Разрешаем каждую строку: сначала по subdomain (A), затем по ИНН (B).
    row_stands = []
    for sub, inn in zip(subs, inns):
        stand = mapping.get(sub.lower())
        if stand is None and inn:
            inn_clean = inn.strip().replace(chr(39), "").replace(chr(34), "")
            stand = inn_mapping.get(inn_clean)
        row_stands.append(stand)
    stands = list(dict.fromkeys(s for s in row_stands if s))
    resolved = sum(1 for s in row_stands if s)
    print_info(f"Соответствие найдено для {resolved}/{len(subs)} строк ({len(stands)} стендов).")
    unresolved = len(subs) - resolved
    if unresolved:
        print_warn(f"Не найдено в конфигурациях: {unresolved} строк — для них будет пометка.")
    if not stands:
        print_error("Ни одна строка из таблицы не найдена в конфигурациях репозитория (по subdomain и ИНН).")
        return
    sql = _token_utilization_sql(stands, days)
    ch_app_url = ("https://argocd-tb01.cowork.ru/applications/langfuse/langfuse"
                  "?view=tree&node=%2FPod%2Flangfuse%2Fchi-langfuse-default-0-0-0&tab=exec&resource=")
    open_url(ch_app_url)
    print_info("Открыт под chi-langfuse-default-0-0-0 (приложение langfuse, вкладка exec).")
    result_path = tokens_dir / TOKENS_RESULT_FILE
    result_path.write_text("", encoding="utf-8")
    print("Выполните команду в открытом веб-терминале пода:")
    print()
    print(colorize(f"clickhouse-client --query \"{sql}\"", CYAN))
    print()
    print_info(f"Скопируйте весь вывод команды в файл: {result_path}")
    answer = input(f"{action_label()} Заполнили файл? (yes/no, Enter — yes): ").strip().lower()
    if answer in ("n", "no"):
        print_warn("Операция отменена.")
        return
    lines = _read_result_file(result_path)
    if not lines:
        print_warn(f"Файл {result_path} пуст или содержит только служебные строки.")
        answer2 = input_with_timeout(f"{action_label()} Продолжить с пометкой «Данных нет»? (yes/no): ", INPUT_TIMEOUT)
        if answer2 is None or answer2.strip().lower() not in ("y", "yes"):
            print_warn("Операция отменена.")
            return
    totals_by_env = _parse_ch_env_output(lines)
    if not totals_by_env:
        print_warn("Не удалось разобрать результат: ожидается табличка «environment вызовы input output total» (табы или пробелы).")
        print_info("Если данных действительно нет — так и оставьте: в столбце D будет пометка «Данных нет».")
    else:
        matched = sum(1 for s in stands if s in totals_by_env)
        print_info(f"Разобрано результатов по {len(totals_by_env)} стендам, из запрошенных {len(stands)} найдено {matched}.")
    # Опциональный шаг: подсчёт пользователей на каждом стенде через heimdall-auth.
    # Считаем только для стендов, где были вызовы LLM (есть данные в ClickHouse).
    users_by_env = {}
    invite_by_env = {}
    active_stands = [s for s in stands if s in totals_by_env]
    if active_stands:
        print_info(f"Подсчёт пользователей и проверка приглашений по {len(active_stands)} стендам с данными...")
        raw_results = _count_users_for_stands(active_stands, stand_zones, tokens_dir)
        for stand, (cnt, invite) in raw_results.items():
            users_by_env[stand] = cnt
            if invite:
                invite_by_env[stand] = invite
        if users_by_env:
            total_users = sum(users_by_env.values())
            print_info(f"Итого пользователей по всем стендам: {_fmt_int(total_users)}.")
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    out_path = tokens_dir / f"Подсчет утилизации_{stamp}.xlsx"
    try:
        filled, no_data, missing, skipped_hidden = _write_token_report_xlsx(src, out_path, row_stands, totals_by_env, days, users_by_env, invite_by_env)
    except Exception as error:
        print_error(f"Не удалось сохранить Excel: {error}")
        return
    print_divider()
    print_info(f"Готово: {filled} строк с аналитикой, {no_data} без данных за период, {missing} не найдено в репозитории.")
    if skipped_hidden:
        print_info(f"Свернутые строки не тронуты: {skipped_hidden}.")
    print_info(f"Файл сохранён: {out_path}")

def _langfuse_ch_query(stand_name, days):
    """SQL для прямого подсчёта токенов в ClickHouse Langfuse (разбивка по моделям + input/output)."""
    return (
        "SELECT provided_model_name AS model, count() AS calls, "
        "sum(usage_details['input']) AS input_tokens, sum(usage_details['output']) AS output_tokens, "
        "sum(usage_details['total']) AS total_tokens "
        "FROM default.observations "
        f"WHERE environment = '{stand_name}' AND type = 'GENERATION' "
        "AND start_time >= now() - INTERVAL {days} DAY "
        "GROUP BY model ORDER BY total_tokens DESC FORMAT TabSeparated"
    ).format(days=days)


def _langfuse_ch_direct_command(stand_name, days):
    """Команда clickhouse-client для выполнения прямо в поде (веб-терминал ArgoCD), без kubectl."""
    sql = _langfuse_ch_query(stand_name, days)
    dq = chr(34)  # двойная кавычка
    return f"clickhouse-client --query {dq}{sql}{dq}"


def _parse_ch_output(raw):
    """Разбор вывода clickhouse-client; None при пустом/некорректном вводе.

    Понимает и TabSeparated, и пробельный формат (табы из веб-терминала ArgoCD превращаются в пробелы):
    последние 4 поля строки — числа (calls, input, output, total), всё остальное — имя модели.
    Возвращает (rows, totals): rows — [(model, calls, inp, out, total)], totals — (calls, inp, out, total)."""
    parsed = []
    for line in raw.splitlines():
        if not line.strip():
            continue
        split = _ch_split(line)
        if split is None:
            continue
        kind, head, nums = split
        metrics = _ch_metrics(nums)
        if metrics is None:
            continue
        model = head if kind == "tab" else " ".join(head[:-4])
        calls, inp, outp, tot = metrics
        parsed.append((model or "(не указано)", calls, inp, outp, tot))
    if not parsed:
        return None
    totals = (
        sum(r[1] for r in parsed),
        sum(r[2] for r in parsed),
        sum(r[3] for r in parsed),
        sum(r[4] for r in parsed),
    )
    return parsed, totals



def _fmt_int(value):
    """Число с разделителем тысяч по-русски (пробел вместо запятой)."""
    return f"{value:,}".replace(",", " ")



def langfuse_tokens_clickhouse(selected, zone, days):
    """Пункт 8, источник ClickHouse: команда → пользователь выполняет в терминале → вставляет результат."""
    stand_name = selected.get("stand_name") or selected.get("folder", "")
    # ClickHouse один на все зоны — ссылка всегда tb01
    ch_app_url = ("https://argocd-tb01.cowork.ru/applications/langfuse/langfuse"
                  "?view=tree&node=%2FPod%2Flangfuse%2Fchi-langfuse-default-0-0-0&tab=exec&resource=")
    print_divider()
    print_info(f"Прямой подсчёт через ClickHouse (стенд {stand_name}, последние {days} дн.).")
    open_url(ch_app_url)
    print_info(f"Открыт под chi-langfuse-default-0-0-0 (приложение langfuse, вкладка exec).")
    print("Выполните команду в открытом веб-терминале пода:")
    print()
    print(colorize(_langfuse_ch_direct_command(stand_name, days), CYAN))
    print()
    while True:
        raw = input_with_timeout(f"{action_label()} Вставьте результат целиком (Ctrl+C — отмена): ", INPUT_TIMEOUT)
        if raw is None:
            print_warn("Таймаут ожидания. Возврат к меню.")
            return
        if is_help_request(raw):
            from giga_finder.helptext import print_llm_help
            print_llm_help()
            continue
        break
    # пустой результат (команда + промпт, без данных) — нормальная ситуация
    stripped = raw.strip()
    if not stripped or all(
        l.strip().startswith(("root@", "#", "$")) or "clickhouse-client" in l or not l.strip()
        for l in stripped.splitlines()
    ):
        subdomain = selected.get("cowork_subdomain") or ""
        stand_label = f"{stand_name} ({subdomain})" if subdomain else stand_name
        print_divider()
        print_info(f"Данных нет: за последние {days} дн. на стенде {stand_label} не найдено вызовов LLM.")
        return
    parsed = _parse_ch_output(raw)
    if parsed is None:
        print_error("Не удалось разобрать результат: ожидается строка «модель вызовы input output total» (табы или пробелы).")
        return
    _, totals = parsed
    calls, inp, outp, total = totals
    print_divider()
    print_info(f"Токены LLM за последние {days} дн.")
    print(f"Вызовов LLM: {_fmt_int(calls)}")
    print(f"Потрачено за период: {_fmt_int(total)} токенов")
    print(f"  запрос к модели (input):   {_fmt_int(inp)}")
    print(f"  ответ модели (output):     {_fmt_int(outp)}")


def prompt_period(cancel_message):
    """1 — 7 дней, 2 — 30 дней, 3 — спросить число. None — отмена."""
    while True:
        print("  Период:")
        print("   1. Последние 7 дней")
        print("   2. Последние 30 дней")
        print("   3. Другой период (дней)")
        print_hint()
        choice = input_with_timeout("Выберите период (1-3): ", INPUT_TIMEOUT)
        if choice is None:
            print_warn(cancel_message)
            return None
        choice = choice.strip()
        if is_help_request(choice):
            from giga_finder.helptext import print_llm_help
            print_llm_help()
            continue
        if choice == "1":
            return 7
        if choice == "2":
            return 30
        if choice == "3":
            raw = input_with_timeout("Число дней (1-365): ", INPUT_TIMEOUT)
            if raw is None:
                print_warn(cancel_message)
                return None
            raw = raw.strip()
            if raw.isdigit() and 1 <= int(raw) <= 365:
                return int(raw)
            print_warn("Неверное число дней, попробуйте снова.")
            continue
        print_warn("Неверный период, попробуйте снова.")


def langfuse_tokens(selected, zone):
    """Пункт 8: сводка по токенам LLM из ClickHouse напрямую (только чтение)."""
    print()
    days = prompt_period("Таймаут ожидания. Возврат к меню.")
    if days is None:
        return
    langfuse_tokens_clickhouse(selected, zone, days)

