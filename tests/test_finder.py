"""Функциональные тесты текущего поведения Finder."""

import io
import pkgutil
import sys
import unittest
import zipfile
from unittest.mock import patch
from contextlib import redirect_stdout
from pathlib import Path

from tests.loader import load, patch_attr

MOD = load()
GOLDEN = Path(__file__).resolve().parent / "golden"
_ORIGINAL = {
    "input_with_timeout": MOD.input_with_timeout,
    "open_url": MOD.open_url,
    "DSN_CACHE": MOD.DSN_CACHE,
    "_count_users_for_stands": MOD._count_users_for_stands,
}


def _restore_runtime():
    for name, value in _ORIGINAL.items():
        patch_attr(MOD, name, value)
    modules = [MOD]
    package = sys.modules.get("giga_finder")
    if package is not None and hasattr(package, "__path__"):
        for _, module_name, _ in pkgutil.walk_packages(package.__path__, "giga_finder."):
            submodule = sys.modules.get(module_name)
            if submodule is not None:
                modules.append(submodule)
    for module in modules:
        if "input" in vars(module):
            delattr(module, "input")

YGG = "postgresql://yggdrasil:secret@yggdrasil-postgres-rw.cowork-tb01-0001.svc.cluster.local:5432/yggdrasil"
HM = "postgresql://heimdall_auth:secret@heimdall-auth-postgres-rw.cowork-tb01-0001.svc.cluster.local:5432/heimdall_auth"
WS_ID = "11111111-1111-1111-1111-111111111111"
EMAIL = "user@example.com"
EXTRA_ID = "22222222-2222-2222-2222-222222222222"
SESSION_ID = "33333333-3333-3333-3333-333333333333"

TB01 = {
    "path": "cce-prod/cowork/clusters/cce-prod-tb01",
    "label": "cce-prod-tb01",
    "url": "https://argocd-tb01.cowork.ru/applications/cowork/{folder}?resource=",
}
SDB = {
    "path": "cce-sdb-tb01/cowork/stands",
    "label": "cce-sdb-tb01-stands",
    "url": "https://argocd.cowork.gigab2b.ru/applications/cowork/{folder}?resource=",
    "stand_prefixes": ("sber-", "cowork"),
}


def _write_config(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


def _repo(root):
    _write_config(root / TB01["path"] / "cowork-tb01-0001" / "config.yaml", """
stand_name: "cowork-tb01-0001"
cowork_subdomain: "Acme"
admin_email: "Admin@Example.com"
cowork_revision: "release/1"
public_hostname: "acme.example.com"
# broken stand
meta:
  expiry_date: '2026-10-19'
  status: trial
  inn: 'AbC12'
""".strip() + "\n")
    _write_config(root / TB01["path"] / "cowork-tb01-0002" / "config.yaml", """
stand_name: "cowork-tb01-0002"
cowork_subdomain: "empty-one"
""".strip() + "\n")
    _write_config(root / TB01["path"] / "cowork-tb01-0003" / "config.yaml", """
stand_name: "cowork-tb01-0003"
cowork_subdomain: "cowork-tb01-0003"
admin_email: "a@b.c"
expiry_date: "2026-01-01"
inn: "111"
""".strip() + "\n")
    _write_config(root / TB01["path"] / "not-a-stand" / "config.yaml", """
stand_name: "not-a-stand"
cowork_subdomain: "ignored"
admin_email: "x@y.z"
expiry_date: "2026-01-01"
inn: "999"
""".strip() + "\n")
    _write_config(root / SDB["path"] / "sber-acme" / "config.yaml", """
stand_name: "sber-acme"
cowork_subdomain: "sberdemo"
admin_email: "sber@example.com"
expiry_date: "2026-12-01"
inn: "5408008559"
status: active
""".strip() + "\n")
    _write_config(root / SDB["path"] / "other-box" / "config.yaml", """
stand_name: "other-box"
cowork_subdomain: "other"
admin_email: "o@o.o"
expiry_date: "2026-12-01"
inn: "000"
""".strip() + "\n")
    _write_config(root / TB01["path"] / "cowork-stand-acme" / "config.yaml", """
stand_name: "stand-acme"
cowork_subdomain: "acme"
admin_email: "a@acme.ru"
expiry_date: "2026-10-01"
inn: "100"
""".strip() + "\n")
    _write_config(root / TB01["path"] / "cowork-stand-inn" / "config.yaml", """
stand_name: "stand-inn"
cowork_subdomain: "inn-only"
admin_email: "i@inn.ru"
expiry_date: "2026-10-01"
inn: "5408008559"
""".strip() + "\n")
    _write_config(root / TB01["path"] / "cowork-stand-hidden" / "config.yaml", """
stand_name: "stand-hidden"
cowork_subdomain: "hidden-co"
admin_email: "h@h.ru"
expiry_date: "2026-10-01"
inn: "200"
""".strip() + "\n")


def _folders(results):
    return {item["folder"] for item in results}


class SearchTests(unittest.TestCase):
    def setUp(self):
        self.root = Path(self.id().replace(".", "_"))
        # изоляция: временный каталог через TemporaryDirectory в каждом тесте
        import tempfile
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        _repo(self.root)
        self.zones = [TB01, SDB]

    def tearDown(self):
        self.tmp.cleanup()

    def _find(self, **criteria):
        query = {"subdomain": None, "email": None, "inn": None, "stand_name": None}
        query.update(criteria)
        return MOD.find_stands(self.root, self.zones, query)

    def test_read_config_fields(self):
        data = MOD.read_config(self.root / TB01["path"] / "cowork-tb01-0001" / "config.yaml")
        self.assertEqual(data["cowork_subdomain"], "Acme")
        self.assertEqual(data["admin_email"], "Admin@Example.com")
        self.assertEqual(data["public_hostname"], "acme.example.com")
        self.assertEqual(data["status"], "trial")
        self.assertEqual(data["expiry_date"], "'2026-10-19'")
        self.assertEqual(data["inn"], "'AbC12'")
        self.assertTrue(data["broken"])

    def test_substring_email_is_case_insensitive(self):
        found = _folders(self._find(email="admin@example.com"))
        self.assertIn("cowork-tb01-0001", found)

    def test_inn_is_case_sensitive(self):
        self.assertIn("cowork-tb01-0001", _folders(self._find(inn="AbC12")))
        self.assertNotIn("cowork-tb01-0001", _folders(self._find(inn="abc12")))

    def test_empty_stand_hidden_unless_exact_folder(self):
        self.assertNotIn("cowork-tb01-0002", _folders(self._find(subdomain="empty")))
        self.assertIn("cowork-tb01-0002", _folders(self._find(stand_name="cowork-tb01-0002")))

    def test_default_subdomain_hidden_unless_exact(self):
        self.assertNotIn("cowork-tb01-0003", _folders(self._find(subdomain="cowork")))
        self.assertIn("cowork-tb01-0003", _folders(self._find(subdomain="cowork-tb01-0003")))
        self.assertIn("cowork-tb01-0003", _folders(self._find(email="a@b.c")))

    def test_prefix_and_sdb(self):
        found = _folders(self._find(subdomain="sberdemo"))
        self.assertIn("sber-acme", found)
        self.assertNotIn("other-box", found)
        self.assertNotIn("not-a-stand", found)

    def test_result_carries_zone(self):
        found = self._find(subdomain="sberdemo")
        self.assertEqual(found[0]["stand_dir"], "cce-sdb-tb01-stands")
        self.assertEqual(found[0]["zone_path"], SDB["path"])


class ZoneTests(unittest.TestCase):
    def test_allowed_paths(self):
        self.assertTrue(MOD._zone_path_allowed("cce-prod/cowork/clusters/cce-prod-st02"))
        self.assertTrue(MOD._zone_path_allowed("cce-sdb-tb01/cowork/stands"))
        self.assertFalse(MOD._zone_path_allowed("foreign/cowork/clusters/x"))

    def test_resolve_zone_merges_extra(self):
        zone = MOD.resolve_zone({"stand_dir": "cce-sdb-tb01-stands"})
        self.assertEqual(zone["code"], "sdb-tb01")
        self.assertEqual(zone["argocd_base"], "https://argocd.cowork.gigab2b.ru")
        self.assertEqual(zone["links_app"], "cowork-heimdall")
        self.assertIn("sber-", zone["stand_prefixes"])

    def test_resolve_standard_zone(self):
        zone = MOD.resolve_zone({"stand_dir": "cce-prod-tb01"})
        self.assertEqual(zone["code"], "tb01")
        self.assertNotIn("argocd_base", zone)


class DsnTests(unittest.TestCase):
    def test_extract_dsn(self):
        raw = "YGGDRASIL_DATABASE_URL=postgresql://u:p@host:5432/db"
        self.assertEqual(MOD._extract_dsn(raw), "postgresql://u:p@host:5432/db")
        quoted = '"postgresql://u:p@host:5432/db"'
        self.assertEqual(MOD._extract_dsn(quoted), "postgresql://u:p@host:5432/db")
        self.assertIsNone(MOD._extract_dsn("not a dsn"))

    def test_expand_dns_suffix(self):
        raw = "postgresql://u:p@yggdrasil-postgres-rw:5432/yggdrasil"
        expanded = MOD._expand_dns_suffix(raw, "cowork-st02-0503")
        self.assertEqual(
            expanded,
            "postgresql://u:p@yggdrasil-postgres-rw.cowork-st02-0503.svc.cluster.local:5432/yggdrasil",
        )
        self.assertEqual(MOD._expand_dns_suffix(expanded, "cowork-st02-0503"), expanded)

    def test_unparsed_url_uses_password_template(self):
        with redirect_stdout(io.StringIO()):
            url = MOD._resolve_db_url("", "yggdrasil", "yggdrasil-postgres-rw", "yggdrasil", "cowork-tb01-0001", "YGGDRASIL_DATABASE_URL")
        self.assertIn("<ПАРОЛЬ>", url)
        self.assertIn("cowork-tb01-0001.svc.cluster.local", url)

    def test_cache_skips_password_template(self):
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            cache = MOD.DsnCache(Path(tmp) / "dsn.json")
            with redirect_stdout(io.StringIO()):
                cache.store("stand-a", "note", yggdrasil="postgresql://u:p@h/db")
                cache.store("stand-b", "note", yggdrasil="postgresql://u:<ПАРОЛЬ>@h/db")
            self.assertEqual(cache.entry("stand-a")["yggdrasil"], "postgresql://u:p@h/db")
            self.assertEqual(cache.entry("stand-b"), {})
            reloaded = MOD.DsnCache(Path(tmp) / "dsn.json")
            self.assertEqual(reloaded.entry("stand-a")["yggdrasil"], "postgresql://u:p@h/db")

    def test_prompt_db_urls_asks_once_when_both_cached(self):
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            cache = MOD.DsnCache(Path(tmp) / "dsn.json")
            with redirect_stdout(io.StringIO()):
                cache.store("cowork-tb01-0001", "note", yggdrasil=YGG, heimdall=HM)
            prompts = []

            def fake_input(prompt=""):
                prompts.append(prompt)
                return ""

            def fail_open(url):
                raise AssertionError(f"браузер не должен открываться: {url}")

            patch_attr(MOD, "DSN_CACHE", cache)
            patch_attr(MOD, "input", fake_input)
            patch_attr(MOD, "open_url", fail_open)
            try:
                with redirect_stdout(io.StringIO()):
                    ygg, hm = MOD.prompt_db_urls(
                        {"folder": "cowork-tb01-0001", "stand_name": "cowork-tb01-0001"},
                        {"code": "tb01"},
                    )
            finally:
                patch_attr(MOD, "DSN_CACHE", MOD.DsnCache(MOD.DSN_CACHE_PATH))
            self.assertEqual(ygg, YGG)
            self.assertEqual(hm, HM)
            self.assertEqual(len(prompts), 1)

    def tearDown(self):
        _restore_runtime()


class ScriptTests(unittest.TestCase):
    def _assert_golden(self, name, text):
        placeholders = ("__IDS__", "__EMAIL__", "__DRY__", "__YGG__", "__HM__", "__EMAILS__", "__EXTRA_IDS__", "__SESSION_ID__", "__USERS__", "__DOMAIN__")
        for token in placeholders:
            self.assertNotIn(token, text, token)
        path = GOLDEN / name
        self.assertEqual(path.read_text(encoding="utf-8"), text)

    def test_change_owner_golden_and_dry_run_flag(self):
        dry = MOD.generate_change_owner_script([WS_ID], EMAIL, True, YGG, HM)
        apply = MOD.generate_change_owner_script([WS_ID], EMAIL, False, YGG, HM)
        self.assertEqual(dry.replace("DRY_RUN = True", "DRY_RUN = False"), apply)
        self.assertIn(WS_ID, dry)
        self.assertIn(EMAIL, dry)
        self._assert_golden("change_owner.py", dry)

    def test_mass_add_golden(self):
        text = MOD.generate_mass_add_script([WS_ID], YGG, HM)
        self.assertIn(WS_ID, text)
        self._assert_golden("mass_add.py", text)

    def test_add_user_golden(self):
        dry = MOD.generate_add_user_script([WS_ID], [EMAIL], True, YGG, HM)
        apply = MOD.generate_add_user_script([WS_ID], [EMAIL], False, YGG, HM)
        self.assertEqual(dry.replace("DRY_RUN = True", "DRY_RUN = False"), apply)
        self.assertIn(EMAIL, dry)
        self._assert_golden("add_user.py", dry)

    def test_remove_user_golden(self):
        dry = MOD.generate_remove_user_script([WS_ID], [EMAIL], True, YGG, HM, [EXTRA_ID])
        apply = MOD.generate_remove_user_script([WS_ID], [EMAIL], False, YGG, HM, [EXTRA_ID])
        self.assertEqual(dry.replace("DRY_RUN = True", "DRY_RUN = False"), apply)
        self.assertIn(EXTRA_ID, dry)
        self._assert_golden("remove_user.py", dry)

    def test_list_workspaces_golden(self):
        text = MOD.generate_list_workspaces_script(YGG)
        self.assertIn(YGG, text)
        self._assert_golden("list_workspaces.py", text)

    def test_fix_roles_golden(self):
        dry = MOD.generate_fix_empty_roles_script(True, HM)
        apply = MOD.generate_fix_empty_roles_script(False, HM)
        self.assertEqual(dry.replace("DRY_RUN = True", "DRY_RUN = False"), apply)
        self._assert_golden("fix_empty_roles.py", dry)

    def test_reset_sessions_golden(self):
        dry = MOD.generate_reset_stale_sessions_script(True, YGG, SESSION_ID)
        apply = MOD.generate_reset_stale_sessions_script(False, YGG, SESSION_ID)
        self.assertEqual(dry.replace("DRY_RUN = True", "DRY_RUN = False"), apply)
        self.assertIn(SESSION_ID, dry)
        self._assert_golden("reset_stale_sessions.py", dry)

    def test_create_accounts_golden(self):
        text = MOD.generate_create_accounts_script(["user81"], "21-school.ru")
        self.assertTrue(text.startswith("/app/services/auth/.venv/bin/python <<'PY'\n"))
        self.assertIn("user81", text)
        self.assertIn("21-school.ru", text)
        self._assert_golden("create_accounts.sh", text)


class ClickHouseTests(unittest.TestCase):
    def test_tab_and_space_and_null(self):
        parsed = MOD._parse_ch_output("gpt-4\t2\t10\t\\N\t0\nmy model/name 3 1 2 0\n")
        self.assertIsNotNone(parsed)
        rows, totals = parsed
        self.assertEqual(rows[0][0], "gpt-4")
        self.assertEqual(rows[0][4], 10)  # пустой total = input + output
        self.assertEqual(rows[1][0], "my model/name")
        self.assertEqual(rows[1][4], 3)
        self.assertEqual(totals[0], 5)

    def test_empty_output(self):
        self.assertIsNone(MOD._parse_ch_output(""))
        self.assertIsNone(MOD._parse_ch_output("root@pod:/# clickhouse-client --query\n"))
        self.assertEqual(MOD._parse_ch_env_output([]), {})

    def test_env_output(self):
        result = MOD._parse_ch_env_output(["stand-acme\t2\t10\t5\t0", "stand inn 1 4 6 10"])
        self.assertEqual(result["stand-acme"], (2, 10, 5, 15))
        self.assertEqual(result["stand"], (1, 4, 6, 10))

    def test_sql_lists_stands(self):
        sql = MOD._token_utilization_sql(["stand-acme", "stand-inn"], 7)
        self.assertIn("environment IN ('stand-acme', 'stand-inn')", sql)
        self.assertIn("INTERVAL 7 DAY", sql)


class AccountsTests(unittest.TestCase):
    def test_parse_ansi_pairs(self):
        text = (
            "\x1b[32m[OK]\x1b[0m создан\n"
            "  email:    a@b.c\n"
            "  password: Secret1!\n"
            "  user_id:  abc\n"
            "\x1b[32m[OK]\x1b[0m\n"
            "  email:    d@e.f\n"
            "  password: Other2@\n"
        )
        self.assertEqual(MOD.parse_accounts_output(text), [("a@b.c", "Secret1!"), ("d@e.f", "Other2@")])

    def test_xlsx_roundtrip(self):
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "accounts.xlsx"
            MOD._write_accounts_xlsx(path, [("a@b.c", "p&<1>")])
            with zipfile.ZipFile(path) as zf:
                sheet = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
            self.assertIn("a@b.c", sheet)
            self.assertIn("p&amp;&lt;1&gt;", sheet)
            self.assertIn(">email<", sheet)
            self.assertIn(">password<", sheet)


class MenuTests(unittest.TestCase):
    def tearDown(self):
        _restore_runtime()

    def _run_menu(self, func, selected, answers):
        queue = list(answers)

        def fake_timeout(prompt, timeout):
            if not queue:
                raise AssertionError(f"лишний промпт: {prompt}")
            return queue.pop(0)

        patch_attr(MOD, "input_with_timeout", fake_timeout)
        buffer = io.StringIO()
        with redirect_stdout(buffer):
            func(selected)
        return buffer.getvalue()

    def test_stand_menu_back_and_console(self):
        opened = []
        patch_attr(MOD, "open_url", lambda url: opened.append(url))
        queue = ["3", "0"]

        def fake_timeout(prompt, timeout):
            return queue.pop(0)

        patch_attr(MOD, "input_with_timeout", fake_timeout)
        buffer = io.StringIO()
        with redirect_stdout(buffer):
            MOD.stand_action_menu({"folder": "cowork-tb01-0001", "cowork_subdomain": "acme"}, Path("."))
        text = buffer.getvalue()
        self.assertIn("Выбрать действие", text)
        self.assertEqual(opened, ["https://tenant-console.cowork.ru/"])

    def test_workspace_menu_hides_base_spaces_off_demo(self):
        text = self._run_menu(MOD.workspace_menu, {"cowork_subdomain": "acme", "stand_dir": "cce-prod-tb01"}, ["0"])
        self.assertIn("Сменить владельца пространства", text)
        self.assertNotIn("Добавить базовые пространства", text)

    def test_workspace_menu_shows_base_spaces_on_demo(self):
        text = self._run_menu(MOD.workspace_menu, {"cowork_subdomain": "demo", "stand_dir": "cce-prod-st02"}, ["0"])
        self.assertIn("Добавить базовые пространства", text)


class TokenReportTests(unittest.TestCase):
    def setUp(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest("openpyxl не установлен")
        self.openpyxl = openpyxl
        import tempfile
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name) / "repo"
        self.script_dir = Path(self.tmp.name) / "script"
        self.script_dir.mkdir()
        _repo(self.root)
        tokens = self.script_dir / "tokens"
        tokens.mkdir()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "subdomain"
        ws["B1"] = "inn"
        ws["A2"] = "acme"
        ws["B2"] = ""
        ws["A3"] = "hidden-co"
        ws.row_dimensions[3].hidden = True
        ws["A4"] = ""
        ws["B4"] = "5408008559"
        wb.save(tokens / "tokens.xlsx")
        wb.close()

    def tearDown(self):
        _restore_runtime()
        if hasattr(self, "tmp"):
            self.tmp.cleanup()

    def test_report_skips_hidden_rows_and_writes_column_d(self):
        opened = []
        answers = {"period": "1"}

        def fake_timeout(prompt, timeout):
            return answers["period"]

        def fake_input(prompt=""):
            if "Заполнили" in prompt:
                result = self.script_dir / "tokens" / "результат clickhouse.txt"
                result.write_text("stand-acme\t1200\t10\t5\t15\n", encoding="utf-8")
                return "yes"
            return "skip"

        def fake_users(stands, stand_to_zone, tokens_dir):
            return {"stand-acme": 12}

        patch_attr(MOD, "open_url", lambda url: opened.append(url))
        patch_attr(MOD, "input_with_timeout", fake_timeout)
        patch_attr(MOD, "input", fake_input)
        patch_attr(MOD, "_count_users_for_stands", fake_users)
        buffer = io.StringIO()
        with redirect_stdout(buffer):
            MOD.token_utilization_report(self.root, self.script_dir, [TB01, SDB])
        text = buffer.getvalue()
        self.assertIn("stand-acme", text)
        self.assertIn("stand-inn", text)
        self.assertNotIn("stand-hidden", text)
        self.assertTrue(any("chi-langfuse-default-0-0-0" in url for url in opened))
        saved = list((self.script_dir / "tokens").glob("Подсчет утилизации_*.xlsx"))
        self.assertEqual(len(saved), 1)
        wb = self.openpyxl.load_workbook(saved[0])
        ws = wb.active
        cell = ws.cell(row=2, column=4).value
        self.assertIn("Вызовов LLM: 1 200", cell)
        self.assertIn("Пользователей: 12", cell)
        self.assertIn("Потрачено за период: 15 токенов", cell)
        self.assertIn("запрос к модели (input)", cell)
        self.assertIn("ответ модели (output)", cell)
        self.assertIsNone(ws.cell(row=3, column=4).value)
        self.assertIn("Данных нет", ws.cell(row=4, column=4).value)
        wb.close()


class PeriodTests(unittest.TestCase):
    def tearDown(self):
        _restore_runtime()

    def test_option_three_asks_for_a_number_of_days(self):
        import giga_finder.tokens as tokens
        answers = ["3", "14"]

        def fake_timeout(prompt, timeout):
            return answers.pop(0)

        patch_attr(MOD, "input_with_timeout", fake_timeout)
        with redirect_stdout(io.StringIO()):
            days = tokens.prompt_period("Таймаут ожидания. Операция отменена.")
        self.assertEqual(days, 14)
        self.assertEqual(answers, [])


class HelpTests(unittest.TestCase):
    def test_main_help_does_not_touch_the_repo(self):
        with patch.object(sys, "argv", ["finder-pro.py", "--help"]):
            buffer = io.StringIO()
            with redirect_stdout(buffer):
                MOD.main()
        text = buffer.getvalue()
        self.assertIn("┌", text)
        self.assertIn("exec/", text)
        self.assertIn("accounts/", text)
        self.assertNotIn("Среда:", text)

    def test_token_help_describes_the_real_files(self):
        with patch.object(sys, "argv", ["finder-pro.py", "token", "--help"]):
            buffer = io.StringIO()
            with redirect_stdout(buffer):
                MOD.main()
        text = buffer.getvalue()
        self.assertIn("tokens.xlsx", text)
        self.assertIn("столбце D", text)
        self.assertIn("результат clickhouse.txt", text)
        self.assertIn("результат пользователи.txt", text)
        self.assertNotIn("Среда:", text)

    def test_help_box_is_yellow_on_a_tty(self):
        import giga_finder.ui as ui
        buffer = io.StringIO()
        buffer.isatty = lambda: True
        with redirect_stdout(buffer):
            ui.print_help_box("Подсказка", ["короткая строка"])
        text = buffer.getvalue()
        self.assertIn("\033[93m", text)
        self.assertIn("Подсказка", text)
        self.assertIn("┌", text)

    def test_question_mark_in_workspace_menu_shows_help(self):
        queue = ["?", "0"]

        def fake_timeout(prompt, timeout):
            return queue.pop(0)

        patch_attr(MOD, "input_with_timeout", fake_timeout)
        buffer = io.StringIO()
        try:
            with redirect_stdout(buffer):
                MOD.workspace_menu({"cowork_subdomain": "acme", "stand_dir": "cce-prod-tb01"})
        finally:
            _restore_runtime()
        text = buffer.getvalue()
        self.assertIn("┌", text)
        self.assertIn("change_workspace_owner.py", text)
        self.assertNotIn("Добавить базовые пространства", text)


if __name__ == "__main__":
    unittest.main()
